from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent
GENERATED_DIR = BASE_DIR / "generated"
GENERATED_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# STYLES
# ============================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

SUBHEADER_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7",
)

KPI_FILL = PatternFill(
    "solid",
    fgColor="E2F0D9",
)

TITLE_FILL = PatternFill(
    "solid",
    fgColor="17365D",
)

WHITE_FONT = Font(
    color="FFFFFF",
    bold=True,
    size=11,
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

TITLE_FONT = Font(
    color="FFFFFF",
    bold=True,
    size=16,
)

BOLD_FONT = Font(
    bold=True,
)

THIN_BORDER = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2"),
)

INR_FORMAT = '₹#,##0.00'
PERCENT_FORMAT = '0.00%'
NUMBER_FORMAT = '#,##0.00'
INTEGER_FORMAT = '#,##0'
DATE_FORMAT = 'dd-mm-yyyy'


# ============================================================
# NORMALIZATION
# ============================================================

def _normalize(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).strip().lower()

    text = text.replace("_", " ")
    text = text.replace("-", " ")

    text = re.sub(r"[^a-z0-9% ]+", " ", text)
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def _safe_number(value: Any) -> float:
    if value is None or value == "":
        return 0.0

    if isinstance(value, bool):
        return 0.0

    try:
        text = str(value).strip()
        text = text.replace(",", "")
        text = text.replace("₹", "")
        text = text.replace("%", "")
        return float(text)
    except Exception:
        return 0.0


def _safe_filename(value: str) -> str:
    value = str(value).strip()

    if not value:
        value = "gst_sales_report"

    value = re.sub(
        r'[<>:"/\\|?*]',
        "_",
        value,
    )

    value = re.sub(
        r"\s+",
        "_",
        value,
    )

    return value[:120]


# ============================================================
# COLUMN RESOLUTION
# ============================================================

COLUMN_ALIASES = {
    "invoice_no": [
        "invoice no",
        "invoice number",
        "invoice",
        "bill no",
        "bill number",
        "inv no",
        "inv number",
    ],
    "date": [
        "date",
        "bill date",
        "invoice date",
        "transaction date",
    ],
    "party_name": [
        "party name",
        "customer",
        "customer name",
        "party",
        "client",
        "client name",
        "buyer",
        "buyer name",
    ],
    "party_type": [
        "party type",
        "customer type",
        "type",
        "customer category",
    ],
    "gstin": [
        "gstin",
        "gst number",
        "gst no",
        "gstin number",
    ],
    "party_state": [
        "party state",
        "customer state",
        "state",
        "billing state",
        "shipping state",
    ],
    "inter_state": [
        "inter state",
        "interstate",
        "inter state y n",
        "is inter state",
    ],
    "category": [
        "category",
        "product category",
        "item category",
    ],
    "product": [
        "product",
        "product name",
        "item",
        "item name",
    ],
    "hsn": [
        "hsn",
        "hsn code",
        "hsn sac",
        "hsn sac code",
        "sac",
        "sac code",
    ],
    "qty": [
        "qty",
        "quantity",
        "quantity sold",
        "units",
    ],
    "unit_price": [
        "unit price",
        "price",
        "rate",
        "selling price",
        "sale price",
    ],
    "gross_amount": [
        "gross amount",
        "gross sale amount",
        "gross sales",
        "sale amount",
        "sales amount",
        "total amount",
        "amount",
    ],
    "discount_percent": [
        "discount %",
        "discount percent",
        "discount percentage",
    ],
    "discount_amount": [
        "discount amt",
        "discount amount",
        "discount",
    ],
    "taxable_amount": [
        "taxable amount",
        "taxable value",
        "taxable",
    ],
    "gst_rate": [
        "gst rate",
        "gst %",
        "gst percent",
        "tax rate",
        "gst percentage",
    ],
    "cgst": [
        "cgst",
        "cgst amount",
    ],
    "sgst": [
        "sgst",
        "sgst amount",
    ],
    "igst": [
        "igst",
        "igst amount",
    ],
    "total_gst": [
        "total gst",
        "gst",
        "gst amount",
        "tax",
        "tax amount",
    ],
    "invoice_total": [
        "invoice total",
        "invoice value",
        "total invoice value",
        "net amount",
        "net sales",
        "grand total",
    ],
    "cost": [
        "cost",
        "cost amount",
        "purchase cost",
        "product cost",
        "cost price",
    ],
    "profit": [
        "profit",
        "profit amount",
        "net profit",
    ],
    "payment_mode": [
        "payment mode",
        "payment method",
        "mode of payment",
        "payment type",
    ],
}


# ============================================================
# COLUMN RESOLUTION
# ============================================================

def _resolve_column(
    columns: List[str],
    field: str,
) -> Optional[str]:

    aliases = COLUMN_ALIASES.get(field, [])

    normalized_aliases = {
        _normalize(alias)
        for alias in aliases
    }

    # --------------------------------------------------------
    # 1. EXACT MATCH
    #
    # Example:
    # GST -> GST
    # GSTIN -> GSTIN
    #
    # GST will NOT match GSTIN.
    # --------------------------------------------------------

    for column in columns:

        normalized_column = _normalize(column)

        if normalized_column in normalized_aliases:
            return column

    # --------------------------------------------------------
    # 2. SAFE TOKEN MATCH
    #
    # Only match when the complete tokens are the same.
    # This prevents:
    #
    # GST -> GSTIN
    # Product -> Product Category
    # Discount -> Discount %
    #
    # from being incorrectly matched.
    # --------------------------------------------------------

    for column in columns:

        normalized_column = _normalize(column)

        column_tokens = set(
            normalized_column.split()
        )

        if not column_tokens:
            continue

        for alias in normalized_aliases:

            alias_tokens = set(
                alias.split()
            )

            if not alias_tokens:
                continue

            if column_tokens == alias_tokens:
                return column

    return None


# ============================================================
# OPTIONAL NUMBER
# ============================================================

def _optional_number(
    row: Dict[str, Any],
    column: Optional[str],
) -> Optional[float]:

    if not column:
        return None

    value = row.get(column)

    if value is None:
        return None

    if isinstance(value, str):

        value = value.strip()

        if not value:
            return None

    return _safe_number(value)


# ============================================================
# DATA NORMALIZATION
# ============================================================

def _normalize_rows(
    rows: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:

    if not rows:
        return []

    # --------------------------------------------------------
    # Collect all source columns
    # --------------------------------------------------------

    columns = []

    for row in rows:

        for column in row.keys():

            if column not in columns:
                columns.append(column)

    # --------------------------------------------------------
    # Resolve source columns
    # --------------------------------------------------------

    resolved = {
        field: _resolve_column(columns, field)
        for field in COLUMN_ALIASES
    }

    normalized_rows = []

    # --------------------------------------------------------
    # Normalize every row
    # --------------------------------------------------------

    for row in rows:

        item = {}

        # ====================================================
        # TEXT FIELDS
        # ====================================================

        item["Invoice No"] = (
            row.get(resolved["invoice_no"])
            if resolved["invoice_no"]
            else ""
        )

        item["Date"] = (
            row.get(resolved["date"])
            if resolved["date"]
            else ""
        )

        item["Party Name"] = (
            row.get(resolved["party_name"])
            if resolved["party_name"]
            else ""
        )

        item["Party Type"] = (
            row.get(resolved["party_type"])
            if resolved["party_type"]
            else ""
        )

        item["GSTIN"] = (
            row.get(resolved["gstin"])
            if resolved["gstin"]
            else ""
        )

        item["Party State"] = (
            row.get(resolved["party_state"])
            if resolved["party_state"]
            else ""
        )

        item["Inter-State"] = (
            row.get(resolved["inter_state"])
            if resolved["inter_state"]
            else ""
        )

        item["Category"] = (
            row.get(resolved["category"])
            if resolved["category"]
            else ""
        )

        item["Product"] = (
            row.get(resolved["product"])
            if resolved["product"]
            else ""
        )

        item["HSN/SAC Code"] = (
            row.get(resolved["hsn"])
            if resolved["hsn"]
            else ""
        )

        item["Payment Mode"] = (
            row.get(resolved["payment_mode"])
            if resolved["payment_mode"]
            else ""
        )

        # ====================================================
        # NUMERIC FIELDS
        #
        # IMPORTANT:
        # Missing source columns now remain None/blank.
        # They are NOT converted to 0.
        # ====================================================

        item["Qty"] = _optional_number(
            row,
            resolved["qty"],
        )

        item["Unit Price"] = _optional_number(
            row,
            resolved["unit_price"],
        )

        item["Gross Amount"] = _optional_number(
            row,
            resolved["gross_amount"],
        )

        item["Discount %"] = _optional_number(
            row,
            resolved["discount_percent"],
        )

        item["Discount Amt"] = _optional_number(
            row,
            resolved["discount_amount"],
        )

        item["Taxable Amount"] = _optional_number(
            row,
            resolved["taxable_amount"],
        )

        item["GST Rate"] = _optional_number(
            row,
            resolved["gst_rate"],
        )

        item["CGST"] = _optional_number(
            row,
            resolved["cgst"],
        )

        item["SGST"] = _optional_number(
            row,
            resolved["sgst"],
        )

        item["IGST"] = _optional_number(
            row,
            resolved["igst"],
        )

        item["Total GST"] = _optional_number(
            row,
            resolved["total_gst"],
        )

        item["Invoice Total"] = _optional_number(
            row,
            resolved["invoice_total"],
        )

        item["Cost"] = _optional_number(
            row,
            resolved["cost"],
        )

        item["Profit"] = _optional_number(
            row,
            resolved["profit"],
        )

        # ====================================================
        # DERIVED VALUES
        # ====================================================

        # ----------------------------------------------------
        # Gross Amount
        # Only calculate when both Qty and Unit Price exist.
        # ----------------------------------------------------

        if item["Gross Amount"] is None:

            if (
                item["Qty"] is not None
                and item["Unit Price"] is not None
            ):

                item["Gross Amount"] = (
                    item["Qty"]
                    * item["Unit Price"]
                )

        # ----------------------------------------------------
        # Discount Amount
        # Only calculate when Discount % actually exists.
        # ----------------------------------------------------

        if item["Discount Amt"] is None:

            if (
                item["Gross Amount"] is not None
                and item["Discount %"] is not None
            ):

                item["Discount Amt"] = (
                    item["Gross Amount"]
                    * item["Discount %"]
                    / 100
                )

        # ----------------------------------------------------
        # Taxable Amount
        #
        # If source does not provide taxable amount,
        # calculate:
        #
        # Gross Amount - Discount Amount
        # ----------------------------------------------------

        if item["Taxable Amount"] is None:

            if item["Gross Amount"] is not None:

                discount_amount = (
                    item["Discount Amt"]
                    if item["Discount Amt"] is not None
                    else 0
                )

                item["Taxable Amount"] = (
                    item["Gross Amount"]
                    - discount_amount
                )

        # ----------------------------------------------------
        # Total GST
        #
        # Only derive Total GST from CGST + SGST + IGST
        # when at least one of those fields actually exists.
        # ----------------------------------------------------

        if item["Total GST"] is None:

            gst_parts = [
                item["CGST"],
                item["SGST"],
                item["IGST"],
            ]

            if any(
                value is not None
                for value in gst_parts
            ):

                item["Total GST"] = sum(
                    value if value is not None else 0
                    for value in gst_parts
                )

        # ----------------------------------------------------
        # Invoice Total
        #
        # If source doesn't provide invoice total,
        # calculate:
        #
        # Taxable Amount + Total GST
        # ----------------------------------------------------

        if item["Invoice Total"] is None:

            if item["Taxable Amount"] is not None:

                total_gst = (
                    item["Total GST"]
                    if item["Total GST"] is not None
                    else 0
                )

                item["Invoice Total"] = (
                    item["Taxable Amount"]
                    + total_gst
                )

        # ----------------------------------------------------
        # Profit
        #
        # Only calculate profit when Cost exists.
        # ----------------------------------------------------

        if item["Profit"] is None:

            if item["Cost"] is not None:

                taxable_amount = (
                    item["Taxable Amount"]
                    if item["Taxable Amount"] is not None
                    else 0
                )

                item["Profit"] = (
                    taxable_amount
                    - item["Cost"]
                )

        # ----------------------------------------------------
        # Add normalized row
        # ----------------------------------------------------

        normalized_rows.append(item)

    return normalized_rows


# ============================================================
# DATA NORMALIZATION
# ============================================================

def _normalize_rows(
    rows: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:

    if not rows:
        return []

    # --------------------------------------------------------
    # Collect all source columns
    # --------------------------------------------------------

    columns = []

    for row in rows:

        for column in row.keys():

            if column not in columns:
                columns.append(column)

    # --------------------------------------------------------
    # Resolve source columns
    # --------------------------------------------------------

    resolved = {
        field: _resolve_column(columns, field)
        for field in COLUMN_ALIASES
    }

    normalized_rows = []

    # --------------------------------------------------------
    # Normalize each row
    # --------------------------------------------------------

    for row in rows:

        item = {}

        # ====================================================
        # TEXT FIELDS
        # ====================================================

        item["Invoice No"] = (
            row.get(resolved["invoice_no"])
            if resolved["invoice_no"]
            else ""
        )

        item["Date"] = (
            row.get(resolved["date"])
            if resolved["date"]
            else ""
        )

        item["Party Name"] = (
            row.get(resolved["party_name"])
            if resolved["party_name"]
            else ""
        )

        item["Party Type"] = (
            row.get(resolved["party_type"])
            if resolved["party_type"]
            else ""
        )

        item["GSTIN"] = (
            row.get(resolved["gstin"])
            if resolved["gstin"]
            else ""
        )

        item["Party State"] = (
            row.get(resolved["party_state"])
            if resolved["party_state"]
            else ""
        )

        item["Inter-State"] = (
            row.get(resolved["inter_state"])
            if resolved["inter_state"]
            else ""
        )

        item["Category"] = (
            row.get(resolved["category"])
            if resolved["category"]
            else ""
        )

        item["Product"] = (
            row.get(resolved["product"])
            if resolved["product"]
            else ""
        )

        item["HSN/SAC Code"] = (
            row.get(resolved["hsn"])
            if resolved["hsn"]
            else ""
        )

        item["Payment Mode"] = (
            row.get(resolved["payment_mode"])
            if resolved["payment_mode"]
            else ""
        )

        # ====================================================
        # NUMERIC FIELDS
        #
        # IMPORTANT:
        # Missing source fields remain None.
        # They are NOT converted to 0.
        # ====================================================

        def optional_number(
            field_name: str,
        ) -> Optional[float]:

            source_column = resolved.get(field_name)

            if not source_column:
                return None

            value = row.get(source_column)

            if value is None:
                return None

            if isinstance(value, str):

                value = value.strip()

                if not value:
                    return None

            return _safe_number(value)

        item["Qty"] = optional_number("qty")

        item["Unit Price"] = optional_number(
            "unit_price"
        )

        item["Gross Amount"] = optional_number(
            "gross_amount"
        )

        item["Discount %"] = optional_number(
            "discount_percent"
        )

        item["Discount Amt"] = optional_number(
            "discount_amount"
        )

        item["Taxable Amount"] = optional_number(
            "taxable_amount"
        )

        item["GST Rate"] = optional_number(
            "gst_rate"
        )

        item["CGST"] = optional_number(
            "cgst"
        )

        item["SGST"] = optional_number(
            "sgst"
        )

        item["IGST"] = optional_number(
            "igst"
        )

        item["Total GST"] = optional_number(
            "total_gst"
        )

        item["Invoice Total"] = optional_number(
            "invoice_total"
        )

        item["Cost"] = optional_number(
            "cost"
        )

        item["Profit"] = optional_number(
            "profit"
        )

        # ====================================================
        # DERIVED VALUES
        # ====================================================

        # ----------------------------------------------------
        # Gross Amount
        #
        # Only calculate if both quantity and unit price
        # actually exist.
        # ----------------------------------------------------

        if item["Gross Amount"] is None:

            if (
                item["Qty"] is not None
                and item["Unit Price"] is not None
            ):

                item["Gross Amount"] = (
                    item["Qty"]
                    * item["Unit Price"]
                )

        # ----------------------------------------------------
        # Discount Amount
        #
        # Only calculate if Discount % actually exists.
        # ----------------------------------------------------

        if item["Discount Amt"] is None:

            if (
                item["Gross Amount"] is not None
                and item["Discount %"] is not None
            ):

                item["Discount Amt"] = (
                    item["Gross Amount"]
                    * item["Discount %"]
                    / 100
                )

        # ----------------------------------------------------
        # Taxable Amount
        #
        # If the source doesn't provide taxable amount,
        # calculate:
        #
        # Gross Amount - Discount Amount
        # ----------------------------------------------------

        if item["Taxable Amount"] is None:

            if item["Gross Amount"] is not None:

                discount_amount = (
                    item["Discount Amt"]
                    if item["Discount Amt"] is not None
                    else 0
                )

                item["Taxable Amount"] = (
                    item["Gross Amount"]
                    - discount_amount
                )

        # ----------------------------------------------------
        # Total GST
        #
        # If GST exists directly in the source,
        # it remains Total GST.
        #
        # Otherwise calculate it from CGST + SGST + IGST
        # only when at least one of those fields exists.
        # ----------------------------------------------------

        if item["Total GST"] is None:

            gst_parts = [
                item["CGST"],
                item["SGST"],
                item["IGST"],
            ]

            if any(
                value is not None
                for value in gst_parts
            ):

                item["Total GST"] = sum(
                    value if value is not None else 0
                    for value in gst_parts
                )

        # ----------------------------------------------------
        # Invoice Total
        #
        # If source doesn't provide invoice total,
        # calculate:
        #
        # Taxable Amount + Total GST
        # ----------------------------------------------------

        if item["Invoice Total"] is None:

            if item["Taxable Amount"] is not None:

                total_gst = (
                    item["Total GST"]
                    if item["Total GST"] is not None
                    else 0
                )

                item["Invoice Total"] = (
                    item["Taxable Amount"]
                    + total_gst
                )

        # ----------------------------------------------------
        # Profit
        #
        # Only calculate when Cost exists.
        # ----------------------------------------------------

        if item["Profit"] is None:

            if item["Cost"] is not None:

                taxable_amount = (
                    item["Taxable Amount"]
                    if item["Taxable Amount"] is not None
                    else 0
                )

                item["Profit"] = (
                    taxable_amount
                    - item["Cost"]
                )

        normalized_rows.append(item)

    return normalized_rows


# ============================================================
# WORKSHEET HELPERS
# ============================================================

def _style_title(
    ws,
    title: str,
    end_column: int,
):
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_column,
    )

    cell = ws.cell(
        row=1,
        column=1,
        value=title,
    )

    cell.fill = TITLE_FILL
    cell.font = TITLE_FONT
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[1].height = 28


def _style_header(ws, row: int = 2):

    for cell in ws[row]:
        if cell.value is None:
            continue

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = THIN_BORDER

    ws.row_dimensions[row].height = 32


def _auto_width(ws):

    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is None:
                continue

            try:
                length = len(str(cell.value))
                max_length = max(
                    max_length,
                    length,
                )
            except Exception:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            35,
        )


def _apply_body_border(ws):

    for row in ws.iter_rows():

        for cell in row:

            if cell.value is not None:
                cell.border = THIN_BORDER
                cell.alignment = Alignment(
                    vertical="center"
                )


def _freeze(ws, cell: str = "A3"):
    ws.freeze_panes = cell


def _add_table(
    ws,
    start_row: int,
    end_row: int,
    end_column: int,
    name: str,
):

    if end_row < start_row:
        return

    ref = (
        f"A{start_row}:"
        f"{get_column_letter(end_column)}"
        f"{end_row}"
    )

    table = Table(
        displayName=name,
        ref=ref,
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    table.tableStyleInfo = style

    ws.add_table(table)


def _format_currency_columns(
    ws,
    columns: List[int],
    start_row: int = 3,
):

    for column in columns:

        for row in range(
            start_row,
            ws.max_row + 1,
        ):

            ws.cell(
                row=row,
                column=column,
            ).number_format = INR_FORMAT


def _format_percent_column(
    ws,
    column: int,
    start_row: int = 3,
):

    for row in range(
        start_row,
        ws.max_row + 1,
    ):

        ws.cell(
            row=row,
            column=column,
        ).number_format = PERCENT_FORMAT


# ============================================================
# SALE INVOICES
# ============================================================

def _build_sale_invoices(
    wb: Workbook,
    rows: List[Dict[str, Any]],
):

    ws = wb.create_sheet(
        "Sale Invoices"
    )

    columns = [
        "Invoice No.",
        "Date",
        "Party Name",
        "Party Type",
        "GSTIN",
        "Party State",
        "Inter-State (Y/N)",
        "Category",
        "Product",
        "HSN/SAC Code",
        "Qty",
        "Unit Price",
        "Gross Amount",
        "Discount %",
        "Discount Amt",
        "Taxable Amount",
        "GST Rate",
        "CGST",
        "SGST",
        "IGST",
        "Total GST",
        "Invoice Total",
        "Cost",
        "Profit",
        "Payment Mode",
    ]

    _style_title(
        ws,
        "Sale Invoices — Raw Register",
        len(columns),
    )

    for index, column in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=2,
            column=index,
            value=column,
        )

    _style_header(ws)
    for row_index, item in enumerate(
        rows,
        start=3,
    ):

        values = [
            item["Invoice No"],
            item["Date"],
            item["Party Name"],
            item["Party Type"],
            item["GSTIN"],
            item["Party State"],
            item["Inter-State"],
            item["Category"],
            item["Product"],
            item["HSN/SAC Code"],
            item["Qty"],
            item["Unit Price"],
            item["Gross Amount"],
            (
                item["Discount %"] / 100
                if item["Discount %"] is not None
                else None
            ),
            item["Discount Amt"],
            item["Taxable Amount"],
            (
                item["GST Rate"] / 100
                if item["GST Rate"] is not None
                else None
            ),
            item["CGST"],
            item["SGST"],
            item["IGST"],
            item["Total GST"],
            item["Invoice Total"],
            item["Cost"],
            item["Profit"],
            item["Payment Mode"],
        ]

        for column_index, value in enumerate(
            values,
            start=1,
        ):

            cell = ws.cell(
                row=row_index,
                column=column_index,
                value=value,
            
            )

            cell.border = THIN_BORDER

    # Currency
    _format_currency_columns(
        ws,
        [
            12,
            13,
            15,
            16,
            18,
            19,
            20,
            21,
            22,
            23,
            24,
        ],
    )

    # Percentage
    _format_percent_column(ws, 14)
    _format_percent_column(ws, 17)

    # Quantity
    for row in range(3, ws.max_row + 1):
        ws.cell(row, 11).number_format = INTEGER_FORMAT

    _freeze(ws)

    _add_table(
        ws,
        2,
        max(2, ws.max_row),
        len(columns),
        "SaleInvoicesTable",
    )

    _auto_width(ws)

    return ws


# ============================================================
# BUSINESS SUMMARY
# ============================================================

def _build_business_summary(
    wb: Workbook,
    rows: List[Dict[str, Any]],
    business_name: str = "Accessories Business",
    gstin: str = "",
    financial_year: str = "FY 2026-27",
):

    ws = wb.create_sheet(
        "Business Summary",
        0,
    )

    _style_title(
        ws,
        "Business Summary — GST Sales Dashboard",
        6,
    )

    ws["A3"] = "Business Name"
    ws["B3"] = business_name

    ws["A4"] = "GSTIN"
    ws["B4"] = gstin or "N/A"

    ws["A5"] = "Financial Year"
    ws["B5"] = financial_year

    for cell in (
        ws["A3"],
        ws["A4"],
        ws["A5"],
    ):
        cell.font = BOLD_FONT
        cell.fill = SUBHEADER_FILL
        cell.border = THIN_BORDER

    for cell in (
        ws["B3"],
        ws["B4"],
        ws["B5"],
    ):
        cell.border = THIN_BORDER

    # --------------------------------------------------------
    # KPI table
    # --------------------------------------------------------

    ws["A7"] = "KPI"
    ws["B7"] = "Value"

    _style_header(ws, 7)

    kpis = [
        (
            "Total Invoices",
            "=COUNTA('Sale Invoices'!A3:A1048576)",
            INTEGER_FORMAT,
        ),
        (
            "Total Quantity Sold",
            "=SUM('Sale Invoices'!K:K)",
            INTEGER_FORMAT,
        ),
        (
            "Gross Sale Amount",
            "=SUM('Sale Invoices'!M:M)",
            INR_FORMAT,
        ),
        (
            "Total Discount Given",
            "=SUM('Sale Invoices'!O:O)",
            INR_FORMAT,
        ),
        (
            "Total Taxable Amount",
            "=SUM('Sale Invoices'!P:P)",
            INR_FORMAT,
        ),
        (
            "CGST Total",
            "=SUM('Sale Invoices'!R:R)",
            INR_FORMAT,
        ),
        (
            "SGST Total",
            "=SUM('Sale Invoices'!S:S)",
            INR_FORMAT,
        ),
        (
            "IGST Total",
            "=SUM('Sale Invoices'!T:T)",
            INR_FORMAT,
        ),
        (
            "Total GST Collected",
            "=SUM('Sale Invoices'!U:U)",
            INR_FORMAT,
        ),
        (
            "Total Invoice Value",
            "=SUM('Sale Invoices'!V:V)",
            INR_FORMAT,
        ),
        (
            "Total Cost",
            "=SUM('Sale Invoices'!W:W)",
            INR_FORMAT,
        ),
        (
            "Net Profit",
            "=SUM('Sale Invoices'!X:X)",
            INR_FORMAT,
        ),
        (
            "Profit Margin %",
            '=IFERROR(B19/B17,0)',
            PERCENT_FORMAT,
        ),
        (
            "Average Invoice Value",
            '=IFERROR(B17/B8,0)',
            INR_FORMAT,
        ),
        (
            "Number of Parties",
            '=SUMPRODUCT((\'Party-wise Ledger\'!A3:A1048576<>"")/COUNTIF(\'Party-wise Ledger\'!A3:A1048576,\'Party-wise Ledger\'!A3:A1048576&""))',
            INTEGER_FORMAT,
        ),
    ]

    for index, (
        name,
        formula,
        number_format,
    ) in enumerate(
        kpis,
        start=8,
    ):

        ws.cell(
            row=index,
            column=1,
            value=name,
        )

        ws.cell(
            row=index,
            column=2,
            value=formula,
        )

        ws.cell(
            row=index,
            column=1,
        ).fill = KPI_FILL

        ws.cell(
            row=index,
            column=1,
        ).font = BOLD_FONT

        ws.cell(
            row=index,
            column=2,
        ).number_format = number_format

        ws.cell(
            row=index,
            column=1,
        ).border = THIN_BORDER

        ws.cell(
            row=index,
            column=2,
        ).border = THIN_BORDER

    _freeze(ws, "A8")

    _auto_width(ws)

    return ws


# ============================================================
# GST SUMMARY
# ============================================================

def _build_gst_summary(
    wb: Workbook,
    rows: List[Dict[str, Any]],
):

    ws = wb.create_sheet(
        "GST Summary"
    )

    columns = [
        "GST Rate",
        "Taxable Value",
        "CGST",
        "SGST",
        "IGST",
        "Total GST",
        "Invoice Value",
    ]

    _style_title(
        ws,
        "GST Summary — Rate-wise",
        len(columns),
    )

    for index, column in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=2,
            column=index,
            value=column,
        )

    _style_header(ws)

    # ========================================================
    # GST RATE HANDLING
    # ========================================================
    #
    # Use GST Rate ONLY when it is actually available.
    #
    # DO NOT derive GST Rate from:
    #
    #     GST / Taxable Amount
    #
    # because that can produce misleading rates when the
    # uploaded dataset does not contain an actual GST Rate.
    #
    # If GST Rate is missing, create one blank-rate summary
    # row so that GST totals are still visible.
    # ========================================================

    actual_rates = sorted(
        {
            row.get("GST Rate")
            for row in rows
            if row.get("GST Rate") is not None
        }
    )

    has_missing_rate = any(
        row.get("GST Rate") is None
        for row in rows
    )

    summary_rows = []

    # --------------------------------------------------------
    # Actual GST rates
    # --------------------------------------------------------

    for rate in actual_rates:
        summary_rows.append(
            {
                "rate": rate,
                "missing_rate": False,
            }
        )

    # --------------------------------------------------------
    # Missing GST rate
    # --------------------------------------------------------
    #
    # Keep the GST Rate cell blank.
    # The formulas below explicitly use "" as the SUMIF
    # criteria so blank GST Rate cells from Sale Invoices
    # are included.
    # --------------------------------------------------------

    if has_missing_rate:
        summary_rows.append(
            {
                "rate": None,
                "missing_rate": True,
            }
        )

    # --------------------------------------------------------
    # Build summary rows
    # --------------------------------------------------------

    for row_index, summary in enumerate(
        summary_rows,
        start=3,
    ):

        rate = summary["rate"]
        missing_rate = summary["missing_rate"]

        # GST Rate
        if rate is not None:
            ws.cell(
                row=row_index,
                column=1,
                value=rate / 100,
            )
        else:
            # Keep missing GST rate genuinely blank.
            ws.cell(
                row=row_index,
                column=1,
                value=None,
            )

        # ----------------------------------------------------
        # SUMIF criteria
        # ----------------------------------------------------

        if missing_rate:

            criteria = '""'

        else:

            criteria = f"A{row_index}"

        # ----------------------------------------------------
        # Taxable Value
        # Sale Invoices column P
        # ----------------------------------------------------

        ws.cell(
            row=row_index,
            column=2,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!Q:Q,"
                f"{criteria},"
                f"'Sale Invoices'!P:P)"
            ),
        )

        # ----------------------------------------------------
        # CGST
        # Sale Invoices column R
        # ----------------------------------------------------

        ws.cell(
            row=row_index,
            column=3,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!Q:Q,"
                f"{criteria},"
                f"'Sale Invoices'!R:R)"
            ),
        )

        # ----------------------------------------------------
        # SGST
        # Sale Invoices column S
        # ----------------------------------------------------

        ws.cell(
            row=row_index,
            column=4,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!Q:Q,"
                f"{criteria},"
                f"'Sale Invoices'!S:S)"
            ),
        )

        # ----------------------------------------------------
        # IGST
        # Sale Invoices column T
        # ----------------------------------------------------

        ws.cell(
            row=row_index,
            column=5,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!Q:Q,"
                f"{criteria},"
                f"'Sale Invoices'!T:T)"
            ),
        )

        # ----------------------------------------------------
        # Total GST
        # Sale Invoices column U
        # ----------------------------------------------------

        ws.cell(
            row=row_index,
            column=6,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!Q:Q,"
                f"{criteria},"
                f"'Sale Invoices'!U:U)"
            ),
        )

        # ----------------------------------------------------
        # Invoice Value
        # Sale Invoices column V
        # ----------------------------------------------------

        ws.cell(
            row=row_index,
            column=7,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!Q:Q,"
                f"{criteria},"
                f"'Sale Invoices'!V:V)"
            ),
        )

    # ========================================================
    # FORMATTING
    # ========================================================

    _format_percent_column(
        ws,
        1,
    )

    _format_currency_columns(
        ws,
        [2, 3, 4, 5, 6, 7],
    )

    _freeze(ws)

    _add_table(
        ws,
        2,
        max(3, ws.max_row),
        len(columns),
        "GSTSummaryTable",
    )

    _auto_width(ws)

    return ws


# ============================================================
# HSN SUMMARY
# ============================================================

def _build_hsn_summary(
    wb: Workbook,
):

    ws = wb.create_sheet(
        "HSN Summary"
    )

    columns = [
        "HSN Code",
        "Category",
        "Total Qty",
        "Taxable Value",
        "Total GST",
        "Total Invoice Value",
    ]

    _style_title(
        ws,
        "HSN Summary",
        len(columns),
    )

    for index, column in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=2,
            column=index,
            value=column,
        )

    _style_header(ws)

    # The formulas use UNIQUE/FILTER where supported by
    # modern Excel. For older Excel, the sheet remains
    # available and can be populated by the next version
    # of the report engine.

    ws["A3"] = (
        '=SORT(UNIQUE(FILTER('
        "'Sale Invoices'!J3:J1048576,"
        "'Sale Invoices'!J3:J1048576<>\"\""
        ')))'
    )

    # The remaining columns use dynamic formulas
    # based on the HSN in column A.
    ws["B3"] = (
        '=IFERROR(INDEX('
        "'Sale Invoices'!H:H,"
        'MATCH(A3,\'Sale Invoices\'!J:J,0)'
        '),"")'
    )

    ws["C3"] = (
        '=SUMIF('
        "'Sale Invoices'!J:J,"
        "A3,"
        "'Sale Invoices'!K:K)"
    )

    ws["D3"] = (
        '=SUMIF('
        "'Sale Invoices'!J:J,"
        "A3,"
        "'Sale Invoices'!P:P)"
    )

    ws["E3"] = (
        '=SUMIF('
        "'Sale Invoices'!J:J,"
        "A3,"
        "'Sale Invoices'!U:U)"
    )

    ws["F3"] = (
        '=SUMIF('
        "'Sale Invoices'!J:J,"
        "A3,"
        "'Sale Invoices'!V:V)"
    )

    _format_currency_columns(
        ws,
        [4, 5, 6],
    )

    _freeze(ws)

    _auto_width(ws)

    return ws


# ============================================================
# PARTY LEDGER
# ============================================================

def _build_party_ledger(
    wb: Workbook,
    rows: List[Dict[str, Any]],
):

    ws = wb.create_sheet(
        "Party-wise Ledger"
    )

    columns = [
        "Party Name",
        "Type",
        "GSTIN",
        "State",
        "Invoice Count",
        "Taxable Value",
        "GST",
        "Invoice Value",
        "Profit",
    ]

    _style_title(
        ws,
        "Party-wise Ledger",
        len(columns),
    )

    for index, column in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=2,
            column=index,
            value=column,
        )

    _style_header(ws)

    parties = []

    seen = set()

    for row in rows:

        party = str(
            row["Party Name"] or ""
        ).strip()

        if not party:
            party = "(Unknown)"

        if party not in seen:
            seen.add(party)
            parties.append(party)

    for row_index, party in enumerate(
        parties,
        start=3,
    ):

        ws.cell(
            row=row_index,
            column=1,
            value=party,
        )

        ws.cell(
            row=row_index,
            column=2,
            value=(
                rows[
                    next(
                        (
                            i
                            for i, item in enumerate(rows)
                            if str(
                                item["Party Name"] or ""
                            ).strip()
                            == party
                        ),
                        0,
                    )
                ]["Party Type"]
                or "Unclassified"
            ),
        )

        ws.cell(
            row=row_index,
            column=3,
            value=(
                rows[
                    next(
                        (
                            i
                            for i, item in enumerate(rows)
                            if str(
                                item["Party Name"] or ""
                            ).strip()
                            == party
                        ),
                        0,
                    )
                ]["GSTIN"]
                or "Unregistered"
            ),
        )

        ws.cell(
            row=row_index,
            column=4,
            value=(
                rows[
                    next(
                        (
                            i
                            for i, item in enumerate(rows)
                            if str(
                                item["Party Name"] or ""
                            ).strip()
                            == party
                        ),
                        0,
                    )
                ]["Party State"]
            ),
        )

        ws.cell(
            row=row_index,
            column=5,
            value=(
                f'=COUNTIF('
                f"'Sale Invoices'!C:C,"
                f"A{row_index})"
            ),
        )

        ws.cell(
            row=row_index,
            column=6,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!C:C,"
                f"A{row_index},"
                f"'Sale Invoices'!P:P)"
            ),
        )

        ws.cell(
            row=row_index,
            column=7,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!C:C,"
                f"A{row_index},"
                f"'Sale Invoices'!U:U)"
            ),
        )

        ws.cell(
            row=row_index,
            column=8,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!C:C,"
                f"A{row_index},"
                f"'Sale Invoices'!V:V)"
            ),
        )

        ws.cell(
            row=row_index,
            column=9,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!C:C,"
                f"A{row_index},"
                f"'Sale Invoices'!X:X)"
            ),
        )

    _format_currency_columns(
        ws,
        [6, 7, 8, 9],
    )

    _freeze(ws)

    _add_table(
        ws,
        2,
        max(2, ws.max_row),
        len(columns),
        "PartyLedgerTable",
    )

    _auto_width(ws)

    return ws


# ============================================================
# CATEGORY SUMMARY
# ============================================================

def _build_category_summary(
    wb: Workbook,
    rows: List[Dict[str, Any]],
):

    ws = wb.create_sheet(
        "Category Summary"
    )

    columns = [
        "Category",
        "HSN",
        "GST Rate",
        "Qty Sold",
        "Taxable Value",
        "GST",
        "Invoice Value",
        "Profit",
        "Margin %",
    ]

    _style_title(
        ws,
        "Category Summary",
        len(columns),
    )

    for index, column in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=2,
            column=index,
            value=column,
        )

    _style_header(ws)

    categories = []

    seen = set()

    for row in rows:

        category = (
            str(row["Category"] or "").strip()
            or "(Unclassified)"
        )

        if category not in seen:
            seen.add(category)
            categories.append(category)

    for row_index, category in enumerate(
        categories,
        start=3,
    ):

        ws.cell(
            row=row_index,
            column=1,
            value=category,
        )

        ws.cell(
            row=row_index,
            column=2,
            value=(
                f'=IFERROR(INDEX('
                f"'Sale Invoices'!J:J,"
                f'MATCH(A{row_index},'
                f"'Sale Invoices'!H:H,0)"
                f'),"")'
            ),
        )

        ws.cell(
            row=row_index,
            column=3,
            value=(
                f'=IFERROR(INDEX('
                f"'Sale Invoices'!Q:Q,"
                f'MATCH(A{row_index},'
                f"'Sale Invoices'!H:H,0)"
                f'),0)'
            ),
        )

        ws.cell(
            row=row_index,
            column=4,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!H:H,"
                f"A{row_index},"
                f"'Sale Invoices'!K:K)"
            ),
        )

        ws.cell(
            row=row_index,
            column=5,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!H:H,"
                f"A{row_index},"
                f"'Sale Invoices'!P:P)"
            ),
        )

        ws.cell(
            row=row_index,
            column=6,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!H:H,"
                f"A{row_index},"
                f"'Sale Invoices'!U:U)"
            ),
        )

        ws.cell(
            row=row_index,
            column=7,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!H:H,"
                f"A{row_index},"
                f"'Sale Invoices'!V:V)"
            ),
        )

        ws.cell(
            row=row_index,
            column=8,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!H:H,"
                f"A{row_index},"
                f"'Sale Invoices'!X:X)"
            ),
        )

        ws.cell(
            row=row_index,
            column=9,
            value=(
                f'=IFERROR(H{row_index}/G{row_index},0)'
            ),
        )

    _format_percent_column(ws, 3)
    _format_currency_columns(
        ws,
        [5, 6, 7, 8],
    )
    _format_percent_column(ws, 9)

    _freeze(ws)

    _add_table(
        ws,
        2,
        max(2, ws.max_row),
        len(columns),
        "CategorySummaryTable",
    )

    _auto_width(ws)

    # --------------------------------------------------------
    # Invoice Value Bar Chart
    # --------------------------------------------------------

    if ws.max_row >= 3:

        chart = BarChart()

        chart.title = "Invoice Value by Category"
        chart.y_axis.title = "Invoice Value (₹)"
        chart.x_axis.title = "Category"

        data = Reference(
            ws,
            min_col=7,
            min_row=2,
            max_row=ws.max_row,
        )

        categories_ref = Reference(
            ws,
            min_col=1,
            min_row=3,
            max_row=ws.max_row,
        )

        chart.add_data(
            data,
            titles_from_data=True,
        )

        chart.set_categories(
            categories_ref
        )

        chart.height = 8
        chart.width = 14

        ws.add_chart(
            chart,
            "K3",
        )

    # --------------------------------------------------------
    # Profit Share Pie Chart
    # --------------------------------------------------------

    if ws.max_row >= 3:

        pie = PieChart()

        pie.title = "Profit Share by Category"

        data = Reference(
            ws,
            min_col=8,
            min_row=2,
            max_row=ws.max_row,
        )

        labels = Reference(
            ws,
            min_col=1,
            min_row=3,
            max_row=ws.max_row,
        )

        pie.add_data(
            data,
            titles_from_data=True,
        )

        pie.set_categories(labels)

        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showLeaderLines = True

        pie.height = 8
        pie.width = 14

        ws.add_chart(
            pie,
            "K20",
        )

    return ws


# ============================================================
# PAYMENT MODE SUMMARY
# ============================================================

def _build_payment_summary(
    wb: Workbook,
    rows: List[Dict[str, Any]],
):

    ws = wb.create_sheet(
        "Payment Mode Summary"
    )

    columns = [
        "Payment Mode",
        "Invoice Count",
        "Invoice Value",
        "% of Total",
    ]

    _style_title(
        ws,
        "Payment Mode Summary",
        len(columns),
    )

    for index, column in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=2,
            column=index,
            value=column,
        )

    _style_header(ws)

    modes = []

    seen = set()

    for row in rows:

        mode = (
            str(row["Payment Mode"] or "").strip()
            or "(Unknown)"
        )

        if mode not in seen:
            seen.add(mode)
            modes.append(mode)

    for row_index, mode in enumerate(
        modes,
        start=3,
    ):

        ws.cell(
            row=row_index,
            column=1,
            value=mode,
        )

        ws.cell(
            row=row_index,
            column=2,
            value=(
                f'=COUNTIF('
                f"'Sale Invoices'!Y:Y,"
                f"A{row_index})"
            ),
        )

        ws.cell(
            row=row_index,
            column=3,
            value=(
                f'=SUMIF('
                f"'Sale Invoices'!Y:Y,"
                f"A{row_index},"
                f"'Sale Invoices'!V:V)"
            ),
        )

        ws.cell(
            row=row_index,
            column=4,
            value=(
                f'=IFERROR('
                f'=IFERROR(C{row_index}/SUM(C:C),0)'
            ),
        )

    _format_currency_columns(
        ws,
        [3],
    )

    _format_percent_column(
        ws,
        4,
    )

    _freeze(ws)

    _add_table(
        ws,
        2,
        max(2, ws.max_row),
        len(columns),
        "PaymentModeSummaryTable",
    )

    _auto_width(ws)

    # Pie chart
    if ws.max_row >= 3:

        pie = PieChart()

        pie.title = "Invoice Value by Payment Mode"

        data = Reference(
            ws,
            min_col=3,
            min_row=2,
            max_row=ws.max_row,
        )

        labels = Reference(
            ws,
            min_col=1,
            min_row=3,
            max_row=ws.max_row,
        )

        pie.add_data(
            data,
            titles_from_data=True,
        )

        pie.set_categories(labels)

        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True

        pie.height = 9
        pie.width = 14

        ws.add_chart(
            pie,
            "F3",
        )

    return ws


# ============================================================
# MONTHLY TREND
# ============================================================

def _build_monthly_trend(
    wb: Workbook,
    rows: List[Dict[str, Any]],
):

    ws = wb.create_sheet(
        "Monthly Trend"
    )

    columns = [
        "Month",
        "Invoices",
        "Taxable Value",
        "GST",
        "Invoice Value",
        "Profit",
    ]

    _style_title(
        ws,
        "Monthly Sales Trend",
        len(columns),
    )

    for index, column in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=2,
            column=index,
            value=column,
        )

    _style_header(ws)

    months = []

    for row in rows:

        date_value = row["Date"]

        if date_value is None:
            continue

        text = str(date_value)

        match = re.match(
            r"(\d{4})[-/](\d{1,2})",
            text,
        )

        if match:

            month = (
                f"{match.group(1)}-"
                f"{int(match.group(2)):02d}"
            )

            if month not in months:
                months.append(month)

    months.sort()

    for row_index, month in enumerate(
        months,
        start=3,
    ):

        year, month_number = month.split("-")

        ws.cell(
            row=row_index,
            column=1,
            value=month,
        )

        # Excel SUMPRODUCT over date text is
        # difficult when source data isn't a real
        # Excel date, therefore these formulas use
        # LEFT on the date column.
        ws.cell(
            row=row_index,
            column=2,
            value=(
                f'=SUMPRODUCT(--('
                f'LEFT(\'Sale Invoices\'!B:B,7)'
                f'="{month}"))'
            ),
        )

        ws.cell(
            row=row_index,
            column=3,
            value=(
                f'=SUMPRODUCT(('
                f'LEFT(\'Sale Invoices\'!B:B,7)'
                f'="{month}")*'
                f"'Sale Invoices'!P:P)"
            ),
        )

        ws.cell(
            row=row_index,
            column=4,
            value=(
                f'=SUMPRODUCT(('
                f'LEFT(\'Sale Invoices\'!B:B,7)'
                f'="{month}")*'
                f"'Sale Invoices'!U:U)"
            ),
        )

        ws.cell(
            row=row_index,
            column=5,
            value=(
                f'=SUMPRODUCT(('
                f'LEFT(\'Sale Invoices\'!B:B,7)'
                f'="{month}")*'
                f"'Sale Invoices'!V:V)"
            ),
        )

        ws.cell(
            row=row_index,
            column=6,
            value=(
                f'=SUMPRODUCT(('
                f'LEFT(\'Sale Invoices\'!B:B,7)'
                f'="{month}")*'
                f"'Sale Invoices'!X:X)"
            ),
        )

    _format_currency_columns(
        ws,
        [3, 4, 5, 6],
    )

    _freeze(ws)

    _add_table(
        ws,
        2,
        max(2, ws.max_row),
        len(columns),
        "MonthlyTrendTable",
    )

    _auto_width(ws)

    # Line chart
    if ws.max_row >= 3:

        chart = LineChart()

        chart.title = "Monthly Invoice Value Trend"
        chart.y_axis.title = "Invoice Value (₹)"
        chart.x_axis.title = "Month"

        data = Reference(
            ws,
            min_col=5,
            min_row=2,
            max_row=ws.max_row,
        )

        categories = Reference(
            ws,
            min_col=1,
            min_row=3,
            max_row=ws.max_row,
        )

        chart.add_data(
            data,
            titles_from_data=True,
        )

        chart.set_categories(
            categories
        )

        chart.height = 9
        chart.width = 16

        ws.add_chart(
            chart,
            "H3",
        )

    return ws


# ============================================================
# MAIN REPORT GENERATOR
# ============================================================

def generate_gst_sales_report(
    rows: List[Dict[str, Any]],
    filename: str = "GST_Sales_Report.xlsx",
    business_name: str = "Accessories Business",
    gstin: str = "",
    financial_year: str = "FY 2026-27",
) -> Dict[str, Any]:

    if not isinstance(rows, list):
        rows = []

    normalized_rows = _normalize_rows(
        rows
    )

    wb = Workbook()

    # Remove default worksheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Build raw register first because
    # summary formulas reference it.
    _build_sale_invoices(
        wb,
        normalized_rows,
    )

    _build_business_summary(
        wb,
        normalized_rows,
        business_name=business_name,
        gstin=gstin,
        financial_year=financial_year,
    )

    _build_gst_summary(
        wb,
        normalized_rows,
    )

    _build_hsn_summary(
        wb,
    )

    _build_party_ledger(
        wb,
        normalized_rows,
    )

    _build_category_summary(
        wb,
        normalized_rows,
    )

    _build_payment_summary(
        wb,
        normalized_rows,
    )

    _build_monthly_trend(
        wb,
        normalized_rows,
    )

    # --------------------------------------------------------
    # Workbook properties
    # --------------------------------------------------------

    wb.properties.title = (
        "Vyapar-style GST Sales Report"
    )

    wb.properties.subject = (
        "GST Sales Report"
    )

    wb.properties.creator = (
        "Dynamic Business AI Agent"
    )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    filename = _safe_filename(
        filename
    )

    if not filename.lower().endswith(
        ".xlsx"
    ):
        filename += ".xlsx"

    output_path = (
        GENERATED_DIR / filename
    )

    counter = 2

    while output_path.exists():

        output_path = (
            GENERATED_DIR
            / (
                f"{Path(filename).stem}_"
                f"{counter}.xlsx"
            )
        )

        counter += 1

    wb.save(output_path)

    return {
        "success": True,
        "format": "xlsx",
        "filename": output_path.name,
        "path": str(output_path),
        "rows": len(normalized_rows),
        "sheets": [
            "Business Summary",
            "Sale Invoices",
            "GST Summary",
            "HSN Summary",
            "Party-wise Ledger",
            "Category Summary",
            "Payment Mode Summary",
            "Monthly Trend",
        ],
    }